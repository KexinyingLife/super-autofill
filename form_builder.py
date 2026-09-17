# -*- coding: utf-8 -*-
"""Non-destructive spreadsheet builder with XML-level in-place editing.

Demonstrates the core technical challenge: xlsxwriter generates native
checkboxes and data-validation dropdowns, but openpyxl.save() strips
them. The solution operates directly on the OOXML zip archive.
"""
from __future__ import annotations

import os
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path
from typing import Any

try:
    import xlsxwriter
except ImportError:
    xlsxwriter = None
try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None

from schema import (
    FIELD_DEFS, FIELD_KEYS, FIELD_HEADS, LIB, CB_DEFAULT, DD_DEFAULT,
    CAT_PRESETS, CAT_FOLLOWUP, default_cells,
)

_MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_REL_NS  = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"

# Vertical spec: (key, kind, label, default, level, tooltip)
SPEC: list[tuple[str, str, str, Any, int, str]] = [
    ("uid",    "meta", "条目ID",   "", 1, ""),
    ("name",   "meta", "条目名称", "", 1, ""),
    ("cat",    "dd",   "类别",     "", 1, "选择后追问生效"),
    ("f1",     "dd",   "字段A(1-5)", "3", 1, "默认3"),
    ("f2",     "dd",   "字段B(1-5)", "3", 1, "默认3"),
    ("f3",     "dd",   "字段C(1-5)", "3", 1, "默认3"),
    ("g1",     "dd",   "· 子项-1", "3:一般", 2, ""),
    ("g2",     "dd",   "· 子项-2", "3:一般", 2, ""),
    ("g3",     "dd",   "· 子项-3", "3:一般", 2, ""),
    ("rec",    "dd",   "是否推荐", "不一定", 1, ""),
    ("recNote","txt",  "推荐理由(≥20字)", "", 1, "空白不提交"),
    ("flag1",  "cb",   "标记A",    False, 1, "勾=是"),
    ("flag2",  "cb",   "标记B",    False, 1, "勾=是"),
    ("sub1",   "dd",   "标签1",    "", 1, ""),
    ("sub2",   "dd",   "标签2(可选)", "", 2, ""),
    ("total",  "dd",   "总分(1-10)", "5", 1, "默认5"),
    ("w1",     "cb",   "意愿1",    True, 1, "勾=是"),
    ("w2",     "cb",   "意愿2",    True, 1, "勾=是"),
    ("w3",     "dd",   "意愿3",    "不一定", 1, "三选"),
    ("status", "meta", "状态",     "", 1, ""),
    ("ts",     "meta", "时间戳",   "", 1, ""),
]

SPEC_KEYS = [s[0] for s in SPEC]

DD_LIB_MAP: dict[str, str] = {
    "cat": "cat", "f1": "s5", "f2": "s5", "f3": "s5",
    "g1": "grade", "g2": "grade", "g3": "grade",
    "rec": "yn", "sub1": "tag", "sub2": "tag",
    "total": "s10", "w3": "yn",
}

BOOL_KEYS = {k for k, _, kind, *_ in SPEC if kind == "cb"}


# ═══════════════════════════════════════════════════════════════════
#  1. Spreadsheet Generation
# ═══════════════════════════════════════════════════════════════════

def _sheet_name(name: str, idx: int) -> str:
    raw = str(name or "item")[:20]
    for ch in "[]:*?/\\":
        raw = raw.replace(ch, "")
    return (raw + "_" + str(idx).zfill(2))[:31]


def _lib_ranges(lib: dict[str, list[str]]) -> dict[str, tuple[int, int]]:
    ranges: dict[str, tuple[int, int]] = {}
    r = 0
    for key, items in lib.items():
        ranges[key] = (r + 2, r + len(items) + 1)
        r += len(items) + 2
    return ranges


def _write_sheet(wb, ws, title: str, cells: dict, lr: dict, lsheet: str) -> None:
    tf = wb.add_format({"bold": True, "font_size": 15, "font_color": "#6A4FA3"})
    l1 = wb.add_format({"font_size": 10.5, "bold": True, "valign": "vcenter", "text_wrap": True})
    l2 = wb.add_format({"font_size": 9, "font_color": "#666", "valign": "vcenter", "text_wrap": True})
    mf = wb.add_format({"font_size": 10, "valign": "vcenter"})
    vf = wb.add_format({"font_size": 11, "valign": "vcenter", "text_wrap": True})
    v2 = wb.add_format({"font_size": 9, "valign": "vcenter", "text_wrap": True})
    wf = wb.add_format({"font_size": 9, "font_color": "#B00020", "valign": "vcenter", "text_wrap": True})

    ws.set_column(0, 0, 36)
    ws.set_column(1, 1, 40)
    ws.write_string(0, 0, title, tf)
    ws.set_row(0, 24)

    for i, (key, kind, label, default, level, tip) in enumerate(SPEC, start=1):
        lf = l1 if level == 1 else l2
        vf_ = vf if level == 1 else v2
        ws.write_string(i, 0, label, lf)
        val = cells.get(key, default)
        if kind == "meta":
            ws.write_string(i, 1, str(val if val is not None else ""), mf)
        elif kind == "cb":
            ws.insert_checkbox(i, 1, bool(val))
        elif kind == "txt":
            ws.write_string(i, 1, str(val or ""), wf if key == "recNote" else vf_)
        else:
            v = val if val not in (None, "") else default
            ws.write_string(i, 1, str(v if v is not None else ""), vf_)
            libkey = DD_LIB_MAP.get(key, key)
            if libkey in lr:
                r1, r2 = lr[libkey]
                ws.data_validation(i, 1, i, 1, {
                    "validate": "list",
                    "source": f"={lsheet}!$A${r1}:$A${r2}",
                })
        if tip and kind != "meta":
            ws.write_comment(i, 0, tip, {"x_scale": 2.0, "y_scale": 1.5})
        ws.set_row(i, 18 if level == 1 else 15)


def make_workbook(rows: list[dict], path: str | Path, lib: dict | None = None) -> None:
    """Generate multi-sheet workbook: one sheet per item, hidden option library."""
    lib = lib or LIB
    wb = xlsxwriter.Workbook(str(path))
    lib_ws = wb.add_worksheet("选项库")
    lr = _lib_ranges(lib)
    r = 0
    for items in lib.values():
        for i, item in enumerate(items, 1):
            lib_ws.write_string(r + i, 0, item)
        r += len(items) + 2
    lib_ws.hide()

    for i, cells in enumerate(rows, 1):
        ws = wb.add_worksheet(_sheet_name(cells.get("name"), i))
        _write_sheet(wb, ws, str(cells.get("name", "")), cells, lr, "选项库")

    ov = wb.add_worksheet("汇总")
    ov.write_string(0, 0, "名称")
    ov.write_string(0, 1, "状态")
    for i, c in enumerate(rows, 1):
        ov.write_string(i, 0, str(c.get("name", "")))
        ov.write_string(i, 1, str(c.get("status", "待处理")))
    ov.set_column(0, 1, 26)
    wb.close()


# ═══════════════════════════════════════════════════════════════════
#  2. Workbook Reading
# ═══════════════════════════════════════════════════════════════════

def _is_vertical(ws) -> bool:
    return str(ws.cell(row=2, column=1).value or "") == "条目ID"


def load_sheets(path: str | Path) -> list[dict]:
    """Read all item sheets from a vertical-layout workbook."""
    wb = load_workbook(str(path), data_only=False)
    out: list[dict] = []
    for ws in wb.worksheets:
        if ws.title in ("选项库", "汇总"):
            continue
        if not _is_vertical(ws):
            continue
        vals: dict[str, Any] = {}
        for i, key in enumerate(FIELD_KEYS, start=1):
            vals[key] = ws.cell(row=i + 1, column=2).value
        if str(vals.get("uid") or "").strip():
            out.append(vals)
    return out


# ═══════════════════════════════════════════════════════════════════
#  3. XML-Level In-Place Editing
# ═══════════════════════════════════════════════════════════════════

def _sheet_map(zf: zipfile.ZipFile) -> dict[str, str]:
    """Sheet name → XML path inside the xlsx zip."""
    wbxml = zf.read("xl/workbook.xml").decode()
    root = ET.fromstring(wbxml)
    rels = ET.fromstring(zf.read("xl/_rels/workbook.xml.rels").decode())
    rid2target = {r.attrib.get("Id"): r.attrib.get("Target") for r in rels}
    result: dict[str, str] = {}
    for el in root.iter():
        if el.tag.endswith("}sheet"):
            name = el.attrib.get("name")
            rid = el.attrib.get("{%s}id" % _REL_NS)
            t = rid2target.get(rid, "").lstrip("/")
            if not t.startswith("xl/"):
                t = "xl/" + t
            result[name] = t
    return result


def _shared_strs(zf: zipfile.ZipFile) -> list[str]:
    try:
        root = ET.fromstring(zf.read("xl/sharedStrings.xml"))
    except KeyError:
        return []
    out: list[str] = []
    for si in root.iter("{%s}si" % _MAIN_NS):
        out.append("".join(t.text or "" for t in si.iter("{%s}t" % _MAIN_NS)))
    return out


def _cell_text(c: ET.Element, shared: list[str]) -> str:
    t = c.attrib.get("t", "")
    v = c.find("{%s}v" % _MAIN_NS)
    if t == "s" and v is not None:
        try:
            return shared[int(v.text or "0")]
        except Exception:
            return ""
    if t == "inlineStr":
        return "".join(x.text or "" for x in c.iter("{%s}t" % _MAIN_NS))
    return (v.text or "") if v is not None else ""


def _set_cell(root: ET.Element, ref: str, value: Any) -> None:
    """Replace a cell's content in-place. Booleans → <c t="b">, strings → inlineStr."""
    target = None
    for row in root.iter("{%s}row" % _MAIN_NS):
        for cell in row.findall("{%s}c" % _MAIN_NS):
            if cell.attrib.get("r") == ref:
                target = cell
                break
        if target is not None:
            break
    if target is None:
        raise ValueError(f"Cell {ref} not found")
    for child in list(target):
        target.remove(child)
    if isinstance(value, bool):
        target.attrib["t"] = "b"
        v_el = ET.SubElement(target, "{%s}v" % _MAIN_NS)
        v_el.text = "1" if value else "0"
    else:
        target.attrib["t"] = "inlineStr"
        is_el = ET.SubElement(target, "{%s}is" % _MAIN_NS)
        t_el = ET.SubElement(is_el, "{%s}t" % _MAIN_NS)
        t_el.text = str(value or "")


def _set_str_xml(xml_str: str, ref: str, text: str) -> str:
    """Set a cell to an inline string (for status row updates)."""
    root = ET.fromstring(xml_str)
    target = None
    for row in root.iter("{%s}row" % _MAIN_NS):
        for c in row.findall("{%s}c" % _MAIN_NS):
            if c.attrib.get("r") == ref:
                target = c
                break
        if target is not None:
            break
    if target is None:
        raise ValueError(f"Cell {ref} not found")
    for child in list(target):
        target.remove(child)
    target.attrib["t"] = "inlineStr"
    is_el = ET.SubElement(target, "{%s}is" % _MAIN_NS)
    t_el = ET.SubElement(is_el, "{%s}t" % _MAIN_NS)
    t_el.text = text or ""
    ET.register_namespace("", _MAIN_NS)
    ET.register_namespace("r", _REL_NS)
    return ET.tostring(root, encoding="unicode", xml_declaration=True)


def _find_item_id(zf: zipfile.ZipFile, sheet_path: str, shared: list[str]) -> str | None:
    """Read the item ID from cell B2 of a sheet."""
    data = zf.read(sheet_path)
    root = ET.fromstring(data)
    for row in root.iter("{%s}row" % _MAIN_NS):
        if int(row.attrib.get("r", "0")) == 2:
            for c in row.findall("{%s}c" % _MAIN_NS):
                if c.attrib.get("r") == "B2":
                    return _cell_text(c, shared)
            break
    return None


def update_answers(path: str | Path, item_id: str, answers: dict[str, Any]) -> int:
    """Write answer values into the spreadsheet via XML surgery.

    Only modifies sheet XML files inside the xlsx zip. Checkboxes,
    dropdowns, styles, and all other content are preserved.
    Returns number of sheets updated.
    """
    allowed = set(FIELD_KEYS) - {"uid", "name", "status", "ts"}
    updates = {k: v for k, v in answers.items() if k in allowed}
    if not updates:
        return 0

    path = str(path)
    with zipfile.ZipFile(path, "r") as zf:
        sm = _sheet_map(zf)
        shared = _shared_strs(zf)
        items: list[tuple[str, bytes]] = []

        for name, target in sm.items():
            if name in ("选项库", "汇总"):
                continue
            sid = _find_item_id(zf, target, shared)
            if str(sid) != str(item_id):
                continue
            data = zf.read(target)
            root = ET.fromstring(data)
            for key, val in updates.items():
                row_idx = SPEC_KEYS.index(key) + 2
                _set_cell(root, f"B{row_idx}", val)
            ET.register_namespace("", _MAIN_NS)
            ET.register_namespace("r", _REL_NS)
            items.append((target, ET.tostring(root, encoding="utf-8", xml_declaration=True)))

    if not items:
        return 0

    tmp = path + ".tmp"
    changed = {t for t, _ in items}
    with zipfile.ZipFile(path, "r") as zin, \
         zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for entry in zin.infolist():
            if entry.filename not in changed:
                zout.writestr(entry, zin.read(entry.filename))
        for target, xml_bytes in items:
            zout.writestr(target, xml_bytes)
    os.replace(tmp, path)
    return len(items)


def update_statuses(path: str | Path, updates: dict[str, tuple[str, str]]) -> int:
    """Update status/timestamp cells without touching other content."""
    s_row = SPEC_KEYS.index("status") + 2
    t_row = SPEC_KEYS.index("ts") + 2
    path = str(path)
    with zipfile.ZipFile(path, "r") as zf:
        sm = _sheet_map(zf)
        shared = _shared_strs(zf)
        items: list[tuple[str, bytes]] = []
        for name, target in sm.items():
            if name in ("选项库", "汇总"):
                continue
            sid = _find_item_id(zf, target, shared)
            if sid and sid in updates:
                status, ts = updates[sid]
                xml_str = zf.read(target).decode("utf-8")
                xml_str = _set_str_xml(xml_str, f"B{s_row}", status)
                xml_str = _set_str_xml(xml_str, f"B{t_row}", ts)
                items.append((target, xml_str.encode("utf-8")))
    if not items:
        return 0
    tmp = path + ".tmp"
    changed = {t for t, _ in items}
    with zipfile.ZipFile(path, "r") as zin, \
         zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for entry in zin.infolist():
            if entry.filename in changed:
                continue
            zout.writestr(entry, zin.read(entry.filename))
        for target, xml_bytes in items:
            zout.writestr(target, xml_bytes)
    os.replace(tmp, path)
    return len(items)
